from __future__ import annotations

from io import BytesIO

from docx import Document

from src.parsers.base import QAParser, requirements_from_rows, requirements_from_text
from src.utils.file_utils import read_bytes


class WordParser(QAParser):
    def parse(self, source, header_mapping=None) -> list:
        try:
            document = Document(BytesIO(read_bytes(source)))
        except Exception as exc:
            raise ValueError(f"Dokumen Word tidak dapat dibaca: {exc}") from exc
        output = requirements_from_text("\n".join(p.text for p in document.paragraphs), "DOCX")
        for table_index, table in enumerate(document.tables, start=1):
            if not table.rows:
                continue
            headers = [cell.text for cell in table.rows[0].cells]
            rows = [dict(zip(headers, [cell.text for cell in row.cells])) for row in table.rows[1:]]
            output.extend(requirements_from_rows(rows, f"tabel DOCX {table_index}", header_mapping))
        return output
